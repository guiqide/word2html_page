#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
演示脚本：创建腾讯游戏法务条款示例并转换为HTML
Demo script: Create a sample Tencent Games legal terms document and convert to HTML
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from word2html import convert_docx_to_html
import os


def create_sample_legal_document():
    """
    创建一个示例法务条款文档
    Create a sample legal terms document
    """
    doc = Document()
    
    # 标题 / Title
    title = doc.add_heading('腾讯游戏用户服务协议', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Tencent Games User Service Agreement')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_format = subtitle.runs[0].font
    subtitle_format.size = Pt(14)
    subtitle_format.color.rgb = RGBColor(128, 128, 128)
    
    # 生效日期 / Effective Date
    date_para = doc.add_paragraph()
    date_para.add_run('生效日期 / Effective Date: 2026年1月22日 / January 22, 2026')
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()  # 空行
    
    # 第一条 / Article 1
    doc.add_heading('第一条 协议的范围和接受', level=1)
    doc.add_heading('Article 1: Scope and Acceptance of Agreement', level=2)
    
    doc.add_paragraph(
        '1.1 本协议是您（以下简称"用户"）与深圳市腾讯计算机系统有限公司'
        '（以下简称"腾讯"）之间关于用户下载、安装、使用腾讯游戏服务的法律协议。'
    )
    doc.add_paragraph(
        '1.1 This agreement is a legal agreement between you (hereinafter referred to as "User") '
        'and Shenzhen Tencent Computer Systems Company Limited (hereinafter referred to as "Tencent") '
        'regarding the user\'s download, installation, and use of Tencent game services.'
    )
    
    doc.add_paragraph(
        '1.2 请您仔细阅读本协议的全部内容。如果您不同意本协议的任何内容，'
        '请不要注册、登录或使用本服务。'
    )
    doc.add_paragraph(
        '1.2 Please read this agreement carefully. If you do not agree with any content of this agreement, '
        'please do not register, log in, or use this service.'
    )
    
    # 第二条 / Article 2
    doc.add_heading('第二条 账号注册与使用', level=1)
    doc.add_heading('Article 2: Account Registration and Use', level=2)
    
    doc.add_paragraph(
        '2.1 用户在使用本服务前需要注册一个游戏账号。游戏账号应当使用手机号码、'
        '电子邮箱或腾讯允许的其他方式注册。'
    )
    doc.add_paragraph(
        '2.1 Users need to register a game account before using this service. '
        'Game accounts should be registered using mobile phone numbers, email addresses, '
        'or other methods permitted by Tencent.'
    )
    
    doc.add_paragraph('2.2 用户承诺：')
    doc.add_paragraph('2.2 Users promise:')
    
    obligations = [
        ('提供真实、准确、完整的注册信息', 'Provide true, accurate, and complete registration information'),
        ('妥善保管账号和密码', 'Properly keep account and password secure'),
        ('不将账号转让或出借他人使用', 'Not transfer or lend the account to others'),
        ('对账号项下的行为承担全部责任', 'Take full responsibility for actions under the account')
    ]
    
    for cn, en in obligations:
        doc.add_paragraph(f'• {cn}', style='List Bullet')
        doc.add_paragraph(f'• {en}', style='List Bullet')
    
    # 第三条 / Article 3
    doc.add_heading('第三条 用户行为规范', level=1)
    doc.add_heading('Article 3: User Conduct Guidelines', level=2)
    
    doc.add_paragraph('3.1 用户在使用本服务时，不得有以下行为：')
    doc.add_paragraph('3.1 When using this service, users shall not engage in the following behaviors:')
    
    prohibited = [
        ('使用外挂、作弊软件或其他非法手段', 'Use cheats, hacking software, or other illegal means'),
        ('发布违法、违规或不当言论', 'Post illegal, violating, or inappropriate comments'),
        ('侵犯他人知识产权或合法权益', 'Infringe on others\' intellectual property or legal rights'),
        ('恶意攻击游戏服务器或破坏游戏秩序', 'Maliciously attack game servers or disrupt game order')
    ]
    
    for i, (cn, en) in enumerate(prohibited, 1):
        doc.add_paragraph(f'{i}. {cn}', style='List Number')
        doc.add_paragraph(f'{i}. {en}', style='List Number')
    
    # 第四条 / Article 4
    doc.add_heading('第四条 服务的变更、中断或终止', level=1)
    doc.add_heading('Article 4: Service Changes, Interruption or Termination', level=2)
    
    doc.add_paragraph(
        '4.1 腾讯有权根据实际情况随时对本服务进行修改、中断或终止，'
        '无需对用户或第三方负责。'
    )
    doc.add_paragraph(
        '4.1 Tencent has the right to modify, interrupt, or terminate this service at any time '
        'according to actual circumstances, without liability to users or third parties.'
    )
    
    # 第五条 / Article 5
    doc.add_heading('第五条 知识产权', level=1)
    doc.add_heading('Article 5: Intellectual Property Rights', level=2)
    
    doc.add_paragraph(
        '5.1 本服务所涉及的游戏软件、图像、音频、视频等内容的知识产权归腾讯所有。'
        '未经腾讯书面许可，任何人不得擅自使用。'
    )
    doc.add_paragraph(
        '5.1 The intellectual property rights of game software, images, audio, video, and other content '
        'involved in this service belong to Tencent. No one may use them without Tencent\'s written permission.'
    )
    
    # 第六条 / Article 6
    doc.add_heading('第六条 隐私保护', level=1)
    doc.add_heading('Article 6: Privacy Protection', level=2)
    
    doc.add_paragraph(
        '6.1 腾讯重视用户的隐私保护。用户在使用本服务时提供的个人信息，'
        '腾讯将按照《腾讯隐私政策》进行收集、使用和保护。'
    )
    doc.add_paragraph(
        '6.1 Tencent values user privacy protection. Personal information provided by users when using '
        'this service will be collected, used, and protected in accordance with the "Tencent Privacy Policy".'
    )
    
    # 第七条 / Article 7
    doc.add_heading('第七条 法律适用与争议解决', level=1)
    doc.add_heading('Article 7: Applicable Law and Dispute Resolution', level=2)
    
    doc.add_paragraph(
        '7.1 本协议的订立、执行和解释及争议的解决均应适用中华人民共和国法律。'
    )
    doc.add_paragraph(
        '7.1 The conclusion, execution, interpretation, and dispute resolution of this agreement '
        'shall be governed by the laws of the People\'s Republic of China.'
    )
    
    doc.add_paragraph(
        '7.2 如双方就本协议内容或其执行发生任何争议，双方应尽力友好协商解决；'
        '协商不成时，任何一方均可向腾讯所在地有管辖权的人民法院提起诉讼。'
    )
    doc.add_paragraph(
        '7.2 If any dispute arises between the parties regarding the content or execution of this agreement, '
        'the parties shall make every effort to resolve it through friendly negotiation; if negotiation fails, '
        'either party may file a lawsuit with the People\'s Court with jurisdiction where Tencent is located.'
    )
    
    # 附则 / Supplementary Provisions
    doc.add_heading('附则', level=1)
    doc.add_heading('Supplementary Provisions', level=2)
    
    doc.add_paragraph(
        '本协议自发布之日起施行。腾讯有权根据需要不时地修订本协议，'
        '并在游戏中公布，无需另行单独通知用户。'
    )
    doc.add_paragraph(
        'This agreement shall come into effect from the date of publication. Tencent has the right '
        'to revise this agreement from time to time as needed and publish it in the game without '
        'separately notifying users.'
    )
    
    # 保存文档
    doc.save('demo_tencent_legal_terms.docx')
    print('✓ 示例文档已创建 / Sample document created: demo_tencent_legal_terms.docx')
    return 'demo_tencent_legal_terms.docx'


def main():
    """主函数 / Main function"""
    print("=" * 70)
    print("腾讯游戏法务条款演示 / Tencent Games Legal Terms Demo")
    print("=" * 70)
    print()
    
    # 创建示例文档
    print("步骤 1: 创建示例法务条款文档...")
    print("Step 1: Creating sample legal terms document...")
    docx_file = create_sample_legal_document()
    print()
    
    # 转换为HTML
    print("步骤 2: 转换为HTML格式...")
    print("Step 2: Converting to HTML format...")
    html_content, messages, output_path = convert_docx_to_html(docx_file)
    print(f'✓ 转换成功 / Conversion successful!')
    print(f'  输出文件 / Output file: {output_path}')
    print()
    
    # 显示文件大小
    docx_size = os.path.getsize(docx_file)
    html_size = os.path.getsize(output_path)
    print(f"文件大小 / File sizes:")
    print(f"  DOCX: {docx_size:,} bytes")
    print(f"  HTML: {html_size:,} bytes")
    print()
    
    # 显示提示
    print("=" * 70)
    print("✓ 演示完成！/ Demo completed!")
    print()
    print("您可以在浏览器中打开HTML文件查看效果：")
    print("You can open the HTML file in a browser to view the result:")
    print(f"  file://{os.path.abspath(output_path)}")
    print("=" * 70)


if __name__ == '__main__':
    main()
