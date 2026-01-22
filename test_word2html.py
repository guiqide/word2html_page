#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Unit tests for word2html converter
"""

import unittest
import os
import tempfile
from pathlib import Path
from docx import Document
from word2html import convert_docx_to_html


class TestWord2HTMLConverter(unittest.TestCase):
    """Test cases for Word to HTML converter"""
    
    def setUp(self):
        """Set up test fixtures"""
        self.temp_dir = tempfile.mkdtemp()
        self.test_docx = os.path.join(self.temp_dir, 'test.docx')
        
        # Create a simple test document
        doc = Document()
        doc.add_heading('Test Heading', 0)
        doc.add_paragraph('This is a test paragraph.')
        doc.add_heading('Subheading', level=1)
        doc.add_paragraph('Another paragraph with content.')
        doc.save(self.test_docx)
    
    def tearDown(self):
        """Clean up test fixtures"""
        import shutil
        if os.path.exists(self.temp_dir):
            shutil.rmtree(self.temp_dir)
    
    def test_basic_conversion(self):
        """Test basic .docx to HTML conversion"""
        output_file = os.path.join(self.temp_dir, 'output.html')
        
        html_content, messages, output_path = convert_docx_to_html(
            self.test_docx, 
            output_file
        )
        
        # Check that output file was created
        self.assertTrue(os.path.exists(output_path))
        
        # Check that HTML content is not empty
        self.assertIsNotNone(html_content)
        self.assertGreater(len(html_content), 0)
        
        # Check that HTML contains expected elements
        self.assertIn('<!DOCTYPE html>', html_content)
        self.assertIn('<html lang="zh-CN">', html_content)
        self.assertIn('Test Heading', html_content)
        self.assertIn('test paragraph', html_content)
    
    def test_default_output_file(self):
        """Test conversion with default output filename"""
        html_content, messages, output_path = convert_docx_to_html(self.test_docx)
        
        # Check that default output file was created
        expected_output = self.test_docx.replace('.docx', '.html')
        self.assertEqual(output_path, expected_output)
        self.assertTrue(os.path.exists(output_path))
    
    def test_nonexistent_file(self):
        """Test handling of non-existent input file"""
        with self.assertRaises(FileNotFoundError):
            convert_docx_to_html('nonexistent.docx')
    
    def test_invalid_format(self):
        """Test handling of non-.docx file"""
        invalid_file = os.path.join(self.temp_dir, 'test.txt')
        with open(invalid_file, 'w') as f:
            f.write('test')
        
        with self.assertRaises(ValueError):
            convert_docx_to_html(invalid_file)
    
    def test_html_structure(self):
        """Test that generated HTML has proper structure"""
        output_file = os.path.join(self.temp_dir, 'output.html')
        html_content, messages, output_path = convert_docx_to_html(
            self.test_docx,
            output_file
        )
        
        # Check for essential HTML elements
        self.assertIn('<head>', html_content)
        self.assertIn('<body>', html_content)
        self.assertIn('<meta charset="UTF-8">', html_content)
        self.assertIn('<style>', html_content)
        
        # Check for Chinese font support
        self.assertIn('Microsoft YaHei', html_content)
    
    def test_output_directory_creation(self):
        """Test that output directories are created if they don't exist"""
        nested_output = os.path.join(self.temp_dir, 'nested', 'dir', 'output.html')
        
        html_content, messages, output_path = convert_docx_to_html(
            self.test_docx,
            nested_output
        )
        
        # Check that nested directory was created
        self.assertTrue(os.path.exists(os.path.dirname(nested_output)))
        self.assertTrue(os.path.exists(output_path))
    
    def test_chinese_content(self):
        """Test conversion of document with Chinese content"""
        chinese_docx = os.path.join(self.temp_dir, 'chinese_test.docx')
        
        doc = Document()
        doc.add_heading('腾讯游戏用户协议', 0)
        doc.add_paragraph('本协议是用户与腾讯之间的协议。')
        doc.save(chinese_docx)
        
        output_file = os.path.join(self.temp_dir, 'chinese_output.html')
        html_content, messages, output_path = convert_docx_to_html(
            chinese_docx,
            output_file
        )
        
        # Check that Chinese content is preserved
        self.assertIn('腾讯游戏用户协议', html_content)
        self.assertIn('本协议是用户与腾讯之间的协议', html_content)


class TestWord2HTMLCLI(unittest.TestCase):
    """Test cases for command-line interface"""
    
    def setUp(self):
        """Set up test fixtures"""
        self.temp_dir = tempfile.mkdtemp()
        self.test_docx = os.path.join(self.temp_dir, 'test.docx')
        
        # Create a simple test document
        doc = Document()
        doc.add_heading('Test Document', 0)
        doc.add_paragraph('Test content.')
        doc.save(self.test_docx)
    
    def tearDown(self):
        """Clean up test fixtures"""
        import shutil
        if os.path.exists(self.temp_dir):
            shutil.rmtree(self.temp_dir)
    
    def test_cli_help(self):
        """Test that CLI help works"""
        import subprocess
        result = subprocess.run(
            ['python', 'word2html.py', '--help'],
            capture_output=True,
            text=True
        )
        self.assertEqual(result.returncode, 0)
        self.assertIn('Word文档转HTML工具', result.stdout)


if __name__ == '__main__':
    unittest.main()
